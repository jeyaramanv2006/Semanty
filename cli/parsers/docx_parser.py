from pathlib import Path
from typing import List
import docx

from cli.parsers.base import BaseParser, ParsedSection


class DocxParser(BaseParser):
    """Word document parser using python-docx with paragraph and heading extraction."""

    def parse(self, file_path: Path) -> List[ParsedSection]:
        doc = docx.Document(str(file_path))
        sections: List[ParsedSection] = []
        
        current_heading = "Document Body"
        buffer: List[str] = []
        
        def flush_buffer(heading_name: str) -> None:
            if buffer:
                combined_text = "\n".join(buffer).strip()
                if combined_text:
                    sections.append(
                        ParsedSection(
                            text=combined_text,
                            location_type="heading",
                            location_value=f"Section: {heading_name}",
                        )
                    )
                buffer.clear()

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name if para.style else ""
            if "Heading" in style_name:
                flush_buffer(current_heading)
                current_heading = text
            else:
                buffer.append(text)

        # Also capture tables if any
        for table_idx, table in enumerate(doc.tables, start=1):
            table_lines: List[str] = []
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    table_lines.append(row_text)
            if table_lines:
                table_content = "\n".join(table_lines)
                buffer.append(f"[Table {table_idx}]\n{table_content}")

        flush_buffer(current_heading)
        return sections
