from pathlib import Path
from typing import List
import pptx

from cli.parsers.base import BaseParser, ParsedSection


class PPTXParser(BaseParser):
    """PowerPoint presentation parser using python-pptx with slide-aware text extraction."""

    def parse(self, file_path: Path) -> List[ParsedSection]:
        prs = pptx.Presentation(str(file_path))
        sections: List[ParsedSection] = []

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_texts: List[str] = []
            
            # Slide title if present
            slide_title = ""
            if slide.shapes.title and slide.shapes.title.text.strip():
                slide_title = slide.shapes.title.text.strip()
                slide_texts.append(f"Title: {slide_title}")

            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_texts.append(text)
                elif shape.has_table:
                    for row in shape.table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            slide_texts.append(row_text)

            # Check speaker notes
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    slide_texts.append(f"Notes: {notes_text}")

            combined_text = "\n".join(slide_texts).strip()
            if combined_text:
                loc_label = f"Slide {slide_idx} ({slide_title})" if slide_title else f"Slide {slide_idx}"
                sections.append(
                    ParsedSection(
                        text=combined_text,
                        location_type="slide",
                        location_value=loc_label,
                    )
                )

        return sections
